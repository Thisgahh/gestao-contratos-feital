import os
import json
import streamlit as st
import pandas as pd
import datetime
import calendar
import smtplib
import ssl
import openpyxl
import hmac
import hashlib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.application import MIMEApplication

# -----------------------------------------------------------------------------
# 0. CONFIGURAÇÃO DE TEMA ESCURO PROFISSIONAL (.streamlit/config.toml)
# -----------------------------------------------------------------------------
def criar_config_toml_automatico():
    pasta_streamlit = ".streamlit"
    caminho_config = os.path.join(pasta_streamlit, "config.toml")
    
    conteudo_config = """[theme]
base = "dark"
primaryColor = "#00E5FF"
backgroundColor = "#0E1117"
secondaryBackgroundColor = "#161B22"
textColor = "#FFFFFF"
font = "sans serif"
"""
    try:
        if not os.path.exists(pasta_streamlit):
            os.makedirs(pasta_streamlit)
        
        with open(caminho_config, "w", encoding="utf-8") as f:
            f.write(conteudo_config)
    except Exception:
        pass

criar_config_toml_automatico()

# -----------------------------------------------------------------------------
# CONFIGURAÇÕES DA PÁGINA E CSS (DARK SLATE & CYAN BRIGHT)
# -----------------------------------------------------------------------------
st.set_page_config(page_title="Gestão de Pagamentos & Contratos", layout="wide")

st.markdown("""
    <style>
    /* 1. ESTRUTURA GERAL */
    html, body, .stApp, [data-testid="stAppViewContainer"], [data-testid="stHeader"] {
        background-color: #0E1117 !important;
        color: #FFFFFF !important;
    }

    /* 2. ABAS SUPERIORES */
    div[data-baseweb="tab-list"] {
        background-color: #161B22 !important;
        padding: 8px 14px !important;
        border-radius: 10px !important;
        border-bottom: 2px solid #00E5FF !important;
    }
    button[data-baseweb="tab"] {
        font-size: 15px !important;
        font-weight: 700 !important;
        color: #94A3B8 !important;
        padding: 10px 22px !important;
    }
    button[data-baseweb="tab"][aria-selected="true"] {
        background-color: #00E5FF !important;
        border-radius: 8px 8px 0px 0px !important;
    }
    button[data-baseweb="tab"][aria-selected="true"] p,
    button[data-baseweb="tab"][aria-selected="true"] span {
        color: #0E1117 !important;
        font-weight: 800 !important;
    }

    /* 3. MENU LATERAL */
    [data-testid="stSidebar"] {
        background-color: #161B22 !important;
        border-right: 1px solid #262C36 !important;
    }

    /* 4. SELECTBOX E INPUTS */
    div[data-baseweb="select"] > div {
        background-color: #161B22 !important;
        color: #FFFFFF !important;
        border: 1px solid #262C36 !important;
        border-radius: 10px !important;
        font-weight: bold !important;
        min-height: 50px !important;
        display: flex !important;
        align-items: center !important;
        padding: 0px 16px !important;
    }
    
    div[data-baseweb="select"] svg {
        fill: #00E5FF !important;
    }
    div[data-baseweb="select"] * {
        color: #FFFFFF !important;
        font-weight: 700 !important;
        font-size: 15px !important;
    }

    /* 5. METRIC CARDS / KPIS */
    [data-testid="stMetric"] {
        background-color: #161B22 !important;
        border: 1px solid #262C36 !important;
        padding: 15px !important;
        border-radius: 10px !important;
        box-shadow: 0 4px 6px rgba(0, 0, 0, 0.3);
        min-height: 80px !important;
    }
    [data-testid="stMetricLabel"] {
        color: #94A3B8 !important;
        font-size: 14px !important;
        font-weight: 600 !important;
    }
    [data-testid="stMetricValue"] {
        color: #00E5FF !important;
        font-weight: 800 !important;
    }

    /* 6. TABELAS (DATA EDITOR) */
    [data-testid="stDataEditor"], [data-testid="stDataFrame"] {
        background-color: #161B22 !important;
        border: 1px solid #262C36 !important;
        border-radius: 8px !important;
    }
    [data-testid="stDataEditor"] [role="columnheader"],
    [data-testid="stDataFrame"] [role="columnheader"] {
        background-color: #1E293B !important;
        border-bottom: 2px solid #00E5FF !important;
    }
    [data-testid="stDataEditor"] [role="columnheader"] *,
    [data-testid="stDataFrame"] [role="columnheader"] * {
        color: #00E5FF !important;
        font-weight: 800 !important;
        text-transform: uppercase !important;
    }
    [data-testid="stDataEditor"] [role="gridcell"],
    [data-testid="stDataFrame"] [role="gridcell"] {
        background-color: #161B22 !important;
        color: #F8FAFC !important;
        border-bottom: 1px solid #21262D !important;
    }

    /* 7. BOTÕES DE AÇÃO */
    div.stButton > button, div.stDownloadButton > button {
        background-color: #00E5FF !important;
        border: none !important;
        border-radius: 6px !important;
        color: #0E1117 !important;
        font-weight: 800 !important;
    }
    div.stButton > button:hover, div.stDownloadButton > button:hover {
        background-color: #00B4D8 !important;
        color: #FFFFFF !important;
    }

    /* 8. INPUTS DE TEXTO */
    input, textarea {
        background-color: #161B22 !important;
        color: #FFFFFF !important;
        border: 1px solid #262C36 !important;
    }
    </style>
""", unsafe_allow_html=True)

# -----------------------------------------------------------------------------
# SISTEMA DE AUTENTICAÇÃO / LOGIN
# -----------------------------------------------------------------------------
USUARIOS_DB = {
    "admin": {
        "nome": "Administrador de TI",
        "senha_hash": hashlib.sha256("admin123".encode()).hexdigest(),
        "perfil": "admin"
    },
    "gestor": {
        "nome": "Gestor de TI",
        "senha_hash": hashlib.sha256("gestor123".encode()).hexdigest(),
        "perfil": "user"
    }
}

def verificar_credenciais(usuario, senha):
    if usuario in USUARIOS_DB:
        senha_hash = hashlib.sha256(senha.encode()).hexdigest()
        if hmac.compare_digest(USUARIOS_DB[usuario]["senha_hash"], senha_hash):
            return USUARIOS_DB[usuario]
    return None

def tela_login():
    st.markdown("<h2 style='text-align: center; color: #00E5FF;'>🔐 Acesso ao Sistema de Pagamentos</h2>", unsafe_allow_html=True)
    st.markdown("<p style='text-align: center; color: #94A3B8;'>Informe suas credenciais corporativas para acessar o painel</p>", unsafe_allow_html=True)
    
    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        with st.form("form_login"):
            usuario_input = st.text_input("Usuário").strip().lower()
            senha_input = st.text_input("Senha", type="password")
            btn_entrar = st.form_submit_button("🔑 Entrar no Sistema", use_container_width=True)
            
            if btn_entrar:
                user_data = verificar_credenciais(usuario_input, senha_input)
                if user_data:
                    st.session_state["logado"] = True
                    st.session_state["usuario_atual"] = user_data["nome"]
                    st.session_state["perfil_atual"] = user_data["perfil"]
                    st.success("✅ Login realizado com sucesso!")
                    st.rerun()
                else:
                    st.error("❌ Usuário ou senha incorretos.")

if "logado" not in st.session_state:
    st.session_state["logado"] = False

if not st.session_state["logado"]:
    tela_login()
    st.stop()

# -----------------------------------------------------------------------------
# CONFIGURAÇÕES DE ARQUIVO E EMAIL
# -----------------------------------------------------------------------------
SMTP_SERVER = "mail.feital.com.br"
SMTP_PORT = 587
EMAIL_REMETENTE = "carlos.rene@feital.com.br"
SENHA_REMETENTE = "S@turno2026"
EMAIL_SETOR_COMPRAS = "carlos.rene@feital.com.br"

EXCEL_PATH = 'CONTRATOS_2026.xlsx'
CONFIG_PATH = 'config_colunas.json'

MESES_LISTA = [
    'JANEIRO', 'FEVEREIRO', 'MARÇO', 'ABRIL', 'MAIO', 'JUNHO',
    'JULHO', 'AGOSTO', 'SETEMBRO', 'OUTUBRO', 'NOVEMBRO', 'DEZEMBRO'
]

LARGURAS_PADRAO = {
    "LOCAL": 110,
    "FATURAMENTO": 170,
    "PRESTADOR": 180,
    "CONTA": 120,
    "SERVIÇO": 320,
    "DT. EMISSÃO NF": 160,
    "VENC.": 110,
    "VALOR_MES": 160,
    "STATUS": 120
}

# Mapping de chaves dos inputs numéricos para o dicionário
INPUT_KEY_MAP = {
    "LOCAL": "inp_local",
    "FATURAMENTO": "inp_fat",
    "PRESTADOR": "inp_prest",
    "CONTA": "inp_conta",
    "SERVIÇO": "inp_serv",
    "DT. EMISSÃO NF": "inp_dt_em",
    "VENC.": "inp_venc",
    "VALOR_MES": "inp_val_mes"
}

# -----------------------------------------------------------------------------
# FUNÇÕES DE SUPORTE
# -----------------------------------------------------------------------------
def carregar_larguras():
    larguras = LARGURAS_PADRAO.copy()
    if os.path.exists(CONFIG_PATH):
        try:
            with open(CONFIG_PATH, 'r', encoding='utf-8') as f:
                dados_salvos = json.load(f)
                larguras.update(dados_salvos)
        except Exception:
            pass
    return larguras

def inicializar_estado_larguras():
    if "larguras" not in st.session_state:
        st.session_state.larguras = carregar_larguras()
    
    # Preenche os estados dos number_input diretamente para nova aba do navegador
    for col_nome, key_input in INPUT_KEY_MAP.items():
        if key_input not in st.session_state:
            st.session_state[key_input] = int(st.session_state.larguras.get(col_nome, LARGURAS_PADRAO.get(col_nome, 120)))

inicializar_estado_larguras()

def salvar_larguras():
    novas_larguras = {}
    for col_nome, key_input in INPUT_KEY_MAP.items():
        val = int(st.session_state.get(key_input, LARGURAS_PADRAO.get(col_nome, 120)))
        novas_larguras[col_nome] = val
        st.session_state.larguras[col_nome] = val

    novas_larguras["STATUS"] = 120
    st.session_state.larguras["STATUS"] = 120

    try:
        with open(CONFIG_PATH, 'w', encoding='utf-8') as f:
            json.dump(novas_larguras, f, indent=4)
        return True
    except Exception as e:
        st.error(f"Erro ao salvar configuração: {e}")
        return False

def formata_br(valor):
    try:
        val_float = float(valor)
        return f"R$ {val_float:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except (ValueError, TypeError):
        return "R$ 0,00"

def converter_para_numero(val):
    if pd.isna(val):
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace('R$', '').replace(' ', '')
    if ',' in s:
        s = s.replace('.', '').replace(',', '.')
    try:
        return float(s)
    except (ValueError, TypeError):
        return 0.0

def calcular_data_valida(dia_param):
    hoje = datetime.date.today()
    if pd.isna(dia_param) or dia_param is None:
        return hoje
    if isinstance(dia_param, (datetime.date, datetime.datetime, pd.Timestamp)):
        return dia_param if isinstance(dia_param, datetime.date) else dia_param.date()
    try:
        dia = int(float(str(dia_param).replace(',', '.')))
        if 1 <= dia <= 31:
            max_dias = calendar.monthrange(hoje.year, hoje.month)[1]
            dia_ajustado = min(dia, max_dias)
            return datetime.date(hoje.year, hoje.month, dia_ajustado)
    except (ValueError, TypeError):
        pass
    return hoje

def enviar_email(destino, assunto, corpo_html, anexo=None):
    try:
        msg = MIMEMultipart()
        msg['From'] = EMAIL_REMETENTE
        msg['To'] = destino
        msg['Subject'] = assunto
        msg.attach(MIMEText(corpo_html, 'html'))

        if anexo is not None:
            part = MIMEApplication(anexo.getvalue(), Name=anexo.name)
            part['Content-Disposition'] = f'attachment; filename="{anexo.name}"'
            msg.attach(part)

        context = ssl.create_default_context()
        with smtplib.SMTP(SMTP_SERVER, SMTP_PORT) as server:
            server.ehlo()
            server.starttls(context=context)
            server.ehlo()
            server.login(EMAIL_REMETENTE, SENHA_REMETENTE)
            server.send_message(msg)
        return True
    except Exception as e:
        st.error(f"Erro ao enviar e-mail: {e}")
        return False

def obter_bytes_excel():
    if os.path.exists(EXCEL_PATH):
        with open(EXCEL_PATH, "rb") as f:
            return f.read()
    return None

def carregar_dados():
    df_consolidado = pd.read_excel(EXCEL_PATH, sheet_name='CONSOLIDADO')
    df_licencas = pd.read_excel(EXCEL_PATH, sheet_name='LICENÇAS_ANUAL')
    
    df_consolidado.columns = [str(c).strip() for c in df_consolidado.columns]
    df_licencas.columns = [str(c).strip() for c in df_licencas.columns]
    
    colunas_texto = ['LOCAL', 'FATURAMENTO', 'PRESTADOR', 'CONTA', 'SERVIÇO']
    for col in colunas_texto:
        if col in df_consolidado.columns:
            df_consolidado[col] = df_consolidado[col].fillna('').astype(str)
        if col in df_licencas.columns:
            df_licencas[col] = df_licencas[col].fillna('').astype(str)

    for col in MESES_LISTA:
        if col in df_consolidado.columns:
            df_consolidado[col] = df_consolidado[col].apply(converter_para_numero)
            
    return df_consolidado, df_licencas

def renderizar_kpis_resumo(df_dash, mes_alvo):
    idx_mes = MESES_LISTA.index(mes_alvo) if mes_alvo in MESES_LISTA else 0
    mes_anterior_nome = MESES_LISTA[idx_mes - 1] if idx_mes > 0 else MESES_LISTA[0]
    
    tot_previsto_mes_ant = df_dash[mes_anterior_nome].sum() if mes_anterior_nome in df_dash.columns else 0.0
    
    val_mes_atual = df_dash[mes_alvo] if mes_alvo in df_dash.columns else pd.Series([0]*len(df_dash))
    df_pagos = df_dash[val_mes_atual > 0]
    df_pendentes = df_dash[val_mes_atual == 0]
    
    qtd_pagos = len(df_pagos)
    qtd_pendentes = len(df_pendentes)
    tot_pago = df_pagos[mes_alvo].sum() if mes_alvo in df_dash.columns else 0.0

    diff = tot_pago - tot_previsto_mes_ant
    str_delta = f"{formata_br(diff)} em relação a {mes_anterior_nome.capitalize()}"

    kpi1, kpi2, kpi3, kpi4 = st.columns(4)
    kpi1.metric(f"Total Previsto (Base {mes_anterior_nome})", formata_br(tot_previsto_mes_ant))
    kpi2.metric(f"Total Confirmado ({mes_alvo})", formata_br(tot_pago), delta=str_delta)
    kpi3.metric("Contratos Concluídos", f"{qtd_pagos} de {len(df_dash)}")
    kpi4.metric("Contratos Pendentes", f"{qtd_pendentes}")

# -----------------------------------------------------------------------------
# INICIALIZAÇÃO DE DADOS
# -----------------------------------------------------------------------------
try:
    if "df_consolidado" not in st.session_state:
        st.session_state.df_consolidado, st.session_state.df_licencas = carregar_dados()
    
    df_consolidado = st.session_state.df_consolidado
    df_licencas = st.session_state.df_licencas
except Exception as e:
    st.error(f"Erro ao carregar o arquivo Excel ({EXCEL_PATH}): {e}")
    st.stop()

# -----------------------------------------------------------------------------
# MENU LATERAL COM SESSÃO DO USUÁRIO
# -----------------------------------------------------------------------------
with st.sidebar:
    st.title("📌 Simples Agenda")
    st.markdown(f"👤 **Usuário:** `{st.session_state.get('usuario_atual')}`")
    st.markdown(f"🛡️ **Perfil:** `{st.session_state.get('perfil_atual').upper()}`")
    
    if st.button("🚪 Sair do Sistema", use_container_width=True):
        st.session_state["logado"] = False
        st.session_state["usuario_atual"] = None
        st.session_state["perfil_atual"] = None
        st.rerun()

    st.markdown("---")
    st.info("💡 Navegue entre o **Dashboard Visão Geral**, as tabelas e o envio de e-mails pelas abas superiores.")

# -----------------------------------------------------------------------------
# CORPO PRINCIPAL
# -----------------------------------------------------------------------------
st.title("📊 Gestão e Controle de Pagamentos")

tab_dashboard, tab_mensais, tab_licencas, tab_solicitacao = st.tabs([
    "📌 Visão Geral (Dashboard)",
    "📑 Contratos Mensais",
    "🔑 Licenças Anuais",
    "📩 Solicitação p/ Compras (E-mail)"
])

idx_tot = df_consolidado[df_consolidado['LOCAL'].astype(str).str.upper() == 'TOTAL'].index
df_dash = df_consolidado.iloc[:idx_tot[0]].copy() if not idx_tot.empty else df_consolidado.copy()

# =============================================================================
# ABA 0: VISÃO GERAL (DASHBOARD)
# =============================================================================
with tab_dashboard:
    mes_atual_nome = MESES_LISTA[datetime.date.today().month - 1]
    
    val_mes_atual = df_dash[mes_atual_nome] if mes_atual_nome in df_dash.columns else pd.Series([0]*len(df_dash))
    qtd_pendentes = len(df_dash[val_mes_atual == 0])

    if qtd_pendentes > 0:
        st.warning(f"⚠️ **Atenção ({mes_atual_nome}):** Existem **{qtd_pendentes} contratos pendentes** de lançamento para este mês.")
    else:
        st.success(f"✅ **Excelente!** Todos os contratos de **{mes_atual_nome}** constam como atualizados.")

    renderizar_kpis_resumo(df_dash, mes_atual_nome)

    st.markdown("---")

    col_g1, col_g2 = st.columns([6, 4])

    with col_g1:
        st.markdown("### 📈 Evolução de Custos Mensais (Visão Anual)")
        
        meses_existentes = [m for m in MESES_LISTA if m in df_dash.columns]
        soma_meses = [df_dash[m].sum() for m in meses_existentes]
        
        df_evolucao = pd.DataFrame({"Mês": meses_existentes, "Total (R$)": soma_meses})
        df_evolucao["Mês"] = pd.Categorical(df_evolucao["Mês"], categories=MESES_LISTA, ordered=True)
        df_evolucao = df_evolucao.sort_values("Mês").set_index("Mês")
        
        st.bar_chart(df_evolucao, color="#00E5FF")

    with col_g2:
        st.markdown("### 🏢 Resumo por Local / Unidade")
        if mes_atual_nome in df_dash.columns:
            df_local = df_dash.groupby('LOCAL')[mes_atual_nome].sum().reset_index()
            df_local.columns = ['Local', 'Total (R$)']
        else:
            df_local = pd.DataFrame(columns=['Local', 'Total (R$)'])
        st.dataframe(
            df_local,
            use_container_width=True,
            hide_index=True,
            column_config={
                "Total (R$)": st.column_config.NumberColumn("Total (R$)", format="R$ %.2f")
            }
        )

# =============================================================================
# ABA 1: CONTRATOS MENSAIS
# =============================================================================
with tab_mensais:
    prestadores_validos = sorted([p for p in df_consolidado['PRESTADOR'].dropna().unique() if str(p).strip().upper() not in ['TOTAL', '']])
    prestadores = ["Todos"] + prestadores_validos
    mes_atual_idx = datetime.date.today().month - 1

    mes_sel_temp = st.session_state.get("sb_mes", MESES_LISTA[mes_atual_idx])

    renderizar_kpis_resumo(df_dash, mes_sel_temp)
    st.markdown("---")

    col_f1, col_f2, col_vazio = st.columns([300, 300, 600])
    with col_f1:
        prestador_sel = st.selectbox("Filtrar por Prestador:", prestadores, key="sb_prestador")
    with col_f2:
        mes_sel = st.selectbox("Selecione o Mês de Referência:", MESES_LISTA, index=mes_atual_idx, key="sb_mes")

    col_exp1, col_exp_vazio = st.columns([500, 700])
    with col_exp1:
        with st.expander("⚙️ Personalizar Largura das Colunas"):
            col_c1, col_c2 = st.columns(2)
            with col_c1:
                st.number_input("LOCAL", step=10, key="inp_local")
                st.number_input("FATURAMENTO", step=10, key="inp_fat")
                st.number_input("PRESTADOR", step=10, key="inp_prest")
                st.number_input("CONTA", step=10, key="inp_conta")
            with col_c2:
                st.number_input("SERVIÇO", step=10, key="inp_serv")
                st.number_input("DT. EMISSÃO NF", step=10, key="inp_dt_em")
                st.number_input("VENC.", step=10, key="inp_venc")
                st.number_input("VALOR DO MÊS", step=10, key="inp_val_mes")

            if st.button("💾 Salvar Layout", use_container_width=True, key="btn_salvar_layout"):
                if salvar_larguras():
                    st.success("✅ Layout salvo com sucesso e gravado em disco!")
                    st.rerun()

    colunas_fixas = ['LOCAL', 'FATURAMENTO', 'PRESTADOR', 'CONTA', 'SERVIÇO', 'DT. EMISSÃO NF', 'VENC.']
    idx_total = df_consolidado[df_consolidado['LOCAL'].astype(str).str.upper() == 'TOTAL'].index
    df_fornecedores = df_consolidado.iloc[:idx_total[0]].copy() if not idx_total.empty else df_consolidado.copy()

    if prestador_sel != "Todos":
        df_fornecedores = df_fornecedores[df_fornecedores['PRESTADOR'] == prestador_sel]

    if mes_sel in df_fornecedores.columns:
        df_fornecedores['STATUS'] = df_fornecedores[mes_sel].apply(lambda v: '✅ OK' if converter_para_numero(v) > 0 else '⏳ Pendente')
    else:
        df_fornecedores[mes_sel] = 0.0
        df_fornecedores['STATUS'] = '⏳ Pendente'

    colunas_exibir = [c for c in colunas_fixas if c in df_fornecedores.columns] + [mes_sel, 'STATUS']

    l_atual = {col: int(st.session_state.get(key_i, st.session_state.larguras.get(col, 120))) for col, key_i in INPUT_KEY_MAP.items()}

    df_editado_mes = st.data_editor(
        df_fornecedores[colunas_exibir],
        use_container_width=False,
        hide_index=True,
        column_config={
            "LOCAL": st.column_config.TextColumn("LOCAL", width=l_atual.get("LOCAL", 110)),
            "FATURAMENTO": st.column_config.TextColumn("FATURAMENTO", width=l_atual.get("FATURAMENTO", 170)),
            "PRESTADOR": st.column_config.TextColumn("PRESTADOR", width=l_atual.get("PRESTADOR", 180)),
            "CONTA": st.column_config.TextColumn("CONTA", width=l_atual.get("CONTA", 120)),
            "SERVIÇO": st.column_config.TextColumn("SERVIÇO", width=l_atual.get("SERVIÇO", 320)),
            "DT. EMISSÃO NF": st.column_config.NumberColumn("DT. EMISSÃO NF", step=1, width=l_atual.get("DT. EMISSÃO NF", 160)),
            "VENC.": st.column_config.NumberColumn("Dia Venc.", step=1, width=l_atual.get("VENC.", 110)),
            mes_sel: st.column_config.NumberColumn(f"Valor ({mes_sel})", format="R$ %.2f", step=0.01, width=l_atual.get("VALOR_MES", 160)),
            "STATUS": st.column_config.TextColumn("STATUS", disabled=True, width=120)
        },
        key=f"editor_mensal_{mes_sel}_{prestador_sel}"
    )

    houve_alteracao = False
    for idx, row in df_editado_mes.iterrows():
        val_atual = converter_para_numero(row[mes_sel])
        val_ant = converter_para_numero(st.session_state.df_consolidado.loc[idx, mes_sel]) if mes_sel in st.session_state.df_consolidado.columns else 0.0
        if val_atual != val_ant:
            st.session_state.df_consolidado.loc[idx, mes_sel] = val_atual
            houve_alteracao = True
    if houve_alteracao:
        st.rerun()

    total_mes = df_editado_mes[mes_sel].apply(converter_para_numero).sum()
    st.markdown(f"### 🧮 SOMA TOTAL ({mes_sel}): **{formata_br(total_mes)}**")

    col_btn1, col_btn2, col_btn_vazio = st.columns([300, 300, 600])
    with col_btn1:
        if st.button("💾 Salvar Planilha em Excel", key="btn_salvar_mensal", use_container_width=True):
            wb = openpyxl.load_workbook(EXCEL_PATH)
            ws = wb['CONSOLIDADO']
            header = [str(cell.value).strip() if cell.value is not None else '' for cell in ws[1]]
            if mes_sel in header:
                col_idx = header.index(mes_sel) + 1
                for idx, row in df_editado_mes.iterrows():
                    ws.cell(row=idx + 2, column=col_idx, value=converter_para_numero(row[mes_sel]))
                for r in range(2, ws.max_row + 1):
                    if str(ws.cell(row=r, column=1).value).strip().upper() == 'TOTAL':
                        ws.cell(row=r, column=col_idx, value=total_mes)
                        break
                wb.save(EXCEL_PATH)
                st.success("✅ Valores salvos com sucesso!")
            else:
                st.error(f"Coluna {mes_sel} não foi encontrada no arquivo Excel.")
    with col_btn2:
        excel_data = obter_bytes_excel()
        if excel_data:
            st.download_button("📥 Baixar Excel", data=excel_data, file_name=EXCEL_PATH, key="btn_download_excel_mensal", use_container_width=True)

# =============================================================================
# ABA 2: LICENÇAS ANUAIS
# =============================================================================
with tab_licencas:
    df_licencas_base = st.session_state.df_licencas.copy()
    mask_total = df_licencas_base['VENCIMENTO'].astype(str).str.upper() == 'TOTAL' if 'VENCIMENTO' in df_licencas_base.columns else pd.Series([False]*len(df_licencas_base))
    df_lic_itens = df_licencas_base[~mask_total].copy() if mask_total.any() else df_licencas_base.copy()

    df_lic_editado = st.data_editor(
        df_lic_itens,
        use_container_width=True,
        hide_index=True,
        column_config={
            "LOCAL": st.column_config.TextColumn("LOCAL", width=l_atual.get("LOCAL", 110)),
            "FATURAMENTO": st.column_config.TextColumn("FATURAMENTO", width=l_atual.get("FATURAMENTO", 170)),
            "PRESTADOR": st.column_config.TextColumn("PRESTADOR", width=l_atual.get("PRESTADOR", 180)),
            "SERVIÇO": st.column_config.TextColumn("SERVIÇO", width=l_atual.get("SERVIÇO", 320)),
            "VALOR PAGO ANO ANTERIOR": st.column_config.NumberColumn("VALOR ANTERIOR", format="R$ %.2f", step=0.01, width=160),
            "VALOR ATUAL": st.column_config.NumberColumn("VALOR ATUAL", format="R$ %.2f", step=0.01, width=160)
        },
        key="editor_licencas"
    )

    col_lic1, col_lic_vazio = st.columns([300, 900])
    with col_lic1:
        if st.button("💾 Salvar Licenças Anuais", key="btn_salvar_licencas", use_container_width=True):
            wb = openpyxl.load_workbook(EXCEL_PATH)
            ws = wb['LICENÇAS_ANUAL']
            for r_idx, row in df_lic_editado.iterrows():
                for c_idx, val in enumerate(row):
                    ws.cell(row=r_idx + 2, column=c_idx + 1, value=val)
            wb.save(EXCEL_PATH)
            st.success("✅ Alterações salvas com sucesso!")

# =============================================================================
# ABA 3: SOLICITAÇÃO (E-MAIL COMPRAS)
# =============================================================================
with tab_solicitacao:
    st.subheader("📩 Enviar Solicitação para Compras")

    def preencher_campos_fornecedor():
        p_selecionado = st.session_state.get("form_prestador_sel")
        if not p_selecionado:
            return

        linha_forn = df_consolidado[df_consolidado['PRESTADOR'] == p_selecionado]
        if not linha_forn.empty:
            dados = linha_forn.iloc[0]
            st.session_state["input_form_desc"] = str(dados.get('SERVIÇO', '')).replace('nan', '').strip()
            
            mes_ref = st.session_state.get('sb_mes', MESES_LISTA[datetime.date.today().month - 1])
            val_mes = converter_para_numero(dados.get(mes_ref, 0)) if mes_ref in dados else 0.0
            if val_mes == 0:
                for m in reversed(MESES_LISTA):
                    if m in dados and converter_para_numero(dados[m]) > 0:
                        val_mes = converter_para_numero(dados[m])
                        break
            st.session_state["input_form_valor"] = float(val_mes)
            st.session_state["input_form_dt_emissao"] = calcular_data_valida(dados.get('DT. EMISSÃO NF', None))
            st.session_state["input_form_dt_venc"] = calcular_data_valida(dados.get('VENC.', None))

    if "input_form_desc" not in st.session_state and prestadores_validos:
        st.session_state["form_prestador_sel"] = prestadores_validos[0]
        preencher_campos_fornecedor()

    with st.container(border=True):
        col_sol_sel, col_sol_vazio = st.columns([300, 900])
        with col_sol_sel:
            st.selectbox(
                "Fornecedor / Prestador",
                prestadores_validos,
                key="form_prestador_sel",
                on_change=preencher_campos_fornecedor
            )

        col_p1, col_p2, col_p3, col_p1_vazio = st.columns([450, 190, 190, 370])
        with col_p1:
            st.text_input("Descrição do Serviço / Item", key="input_form_desc")
        with col_p2:
            st.number_input("Valor da Fatura (R$)", step=10.0, format="%.2f", key="input_form_valor")
        with col_p3:
            st.text_input("Número da NF / Fatura", value="NF-", key="form_nf")

        col_p4, col_p5, col_p2_vazio = st.columns([150, 150, 900])
        with col_p4:
            st.date_input("Data de Emissão da NF", format="DD/MM/YYYY", key="input_form_dt_emissao")
        with col_p5:
            st.date_input("Data de Vencimento", format="DD/MM/YYYY", key="input_form_dt_venc")

        col_obs, col_obs_vazio = st.columns([450, 750])
        with col_obs:
            st.text_area("Observações Adicionais para Compras", value="Segue nota fiscal em anexo para pagamento.", height=100, key="form_obs")

        col_up1, col_up_vazio = st.columns([300, 900])
        with col_up1:
            arquivo_nf = st.file_uploader(
                "📎 Anexar Nota Fiscal / Comprovante (Obrigatório)",
                type=["pdf", "png", "jpg", "jpeg", "xml", "zip"],
                key="form_anexo"
            )

        col_send1, col_send_vazio = st.columns([300, 900])
        with col_send1:
            if st.button("✉️ Enviar E-mail para Compras", key="btn_enviar_solicitacao", use_container_width=True):
                if arquivo_nf is None:
                    st.warning("⚠️ **Atenção:** É obrigatório anexar a Nota Fiscal / Comprovante antes de enviar o e-mail.")
                else:
                    f_nome = st.session_state.form_prestador_sel
                    f_desc = st.session_state.input_form_desc
                    f_val = st.session_state.input_form_valor
                    f_nf = st.session_state.form_nf
                    f_dt_em = st.session_state.input_form_dt_emissao.strftime('%d/%m/%Y')
                    f_dt_vc = st.session_state.input_form_dt_venc.strftime('%d/%m/%Y')
                    f_obs = st.session_state.form_obs

                    corpo_email = f"""
                    <h3>Solicitação de Pedido de Compra</h3>
                    <p>Por gentileza, gerar o Pedido de Compra referente aos dados abaixo:</p>
                    <ul>
                        <li><b>Fornecedor:</b> {f_nome}</li>
                        <li><b>Serviço:</b> {f_desc}</li>
                        <li><b>Número da NF:</b> {f_nf}</li>
                        <li><b>Data de Emissão:</b> {f_dt_em}</li>
                        <li><b>Data de Vencimento:</b> {f_dt_vc}</li>
                        <li><b>Valor:</b> {formata_br(f_val)}</li>
                        <li><b>Observações:</b> {f_obs}</li>
                        <li><b>Anexo:</b> Nota Fiscal enviada em anexo nesta mensagem.</li>
                    </ul>
                    <p><i>E-mail enviado via Sistema de Gestão de Contratos por {st.session_state.get('usuario_atual')}.</i></p>
                    """
                    
                    sucesso = enviar_email(EMAIL_SETOR_COMPRAS, f"Pedido de Compra - {f_nome} (NF {f_nf})", corpo_email, arquivo_nf)
                    
                    if sucesso:
                        st.success(f"✅ E-mail enviado com sucesso para Compras com o anexo **{arquivo_nf.name}** no valor de **{formata_br(f_val)}**!")
